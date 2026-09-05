"""Выгрузка состава актива в Excel и разбор загруженного обратно файла.

Столбцы читаются ПО ЗАГОЛОВКУ, а не по номеру: админ может переставить
колонки местами или спрятать лишние — файл всё равно разберётся. Первый
столбец «id» служебный: по нему строка находит свою запись в базе, поэтому
правки не создают дублей. Пустой id — значит новый человек.
"""
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (ключ в базе, заголовок в файле, ширина)
COLUMNS = (
    ("id",           "id",                    6),
    ("fio",          "ФИО",                   32),
    ("ik_div",       "Направление",           22),
    ("group",        "Группа",                14),
    ("birthday",     "День рождения",         15),
    ("phone",        "Телефон",               16),
    ("email",        "Почта",                 26),
    ("studak",       "Студенческий",          14),
    ("clothes_size", "Размер одежды",         14),
    ("someone_div",  "Другие подразделения",  26),
    ("tg_username",  "Телеграм",              20),
    ("is_active",    "В активе",              10),
    ("registered",   "Привязан телеграм",     18),
)

# Столбцы, которые загрузка игнорирует: они справочные.
READ_ONLY = {"registered"}

TEXT_FIELDS = ("fio", "ik_div", "group", "phone", "email",
               "clothes_size", "someone_div", "tg_username")

# Строка, начинающаяся с этих символов, в Excel считается формулой.
# Часть телефонов в базе лежит как '=+79001234567' (наследие старого
# файла): без пометки «это текст» такая ячейка при обратной загрузке
# вернулась бы пустой и стёрла бы номер. Заодно закрывает формульную
# инъекцию — файл открывают живые люди.
_FORMULA_START = ("=", "+", "-", "@")

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_LOCKED_FILL = PatternFill("solid", fgColor="F2F2F2")


def _norm(value) -> str:
    """Заголовок к сравнимому виду: регистр и пробелы не важны."""
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _cell(key: str, activist, registered: bool):
    if key == "id":
        return activist.id
    if key == "registered":
        return "да" if registered else "нет"
    if key == "is_active":
        return "да" if activist.is_active else "нет"
    if key == "birthday":
        value = activist.birthday
        return value.date() if isinstance(value, datetime) else value
    value = getattr(activist, key, "")
    return "" if value in (None, 0) else value


def build_workbook(activists: list, registered_ids: set[int]) -> io.BytesIO:
    """Файл состава: строка на человека, заголовок закреплён."""
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Актив"

    for col, (_, title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = width

    for row, activist in enumerate(activists, start=2):
        is_registered = activist.id in registered_ids
        for col, (key, _, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row, column=col, value=_cell(key, activist, is_registered))
            if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_START):
                cell.data_type = "s"
            if key == "birthday" and cell.value is not None:
                cell.number_format = "DD.MM.YYYY"
            if key in READ_ONLY or key == "id":
                cell.fill = _LOCKED_FILL

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(2, len(activists) + 1)}"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _to_str(value) -> str:
    """Числа из Excel приезжают как float: 89963479923.0 → '89963479923'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_birthday(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = _to_str(value)
    if not text:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"не понял дату «{text}», нужен вид ДД.ММ.ГГГГ")


def _to_bool(value, default: bool = True) -> bool:
    text = _to_str(value).casefold()
    if not text:
        return default
    return text in ("да", "yes", "true", "1", "+", "х", "x", "v", "да ")


def parse_workbook(data: io.BytesIO) -> tuple[list[dict], list[str]]:
    """Разобрать загруженный файл.

    Возвращает (строки, ошибки). Строки — словари с ключами из COLUMNS плюс
    'row' (номер строки в файле, чтобы человеку было понятно, где чинить).
    Строки с ошибками в результат не попадают: лучше не применить одну, чем
    записать в базу мусор.
    """
    try:
        wb = load_workbook(data, data_only=True)
    except Exception as err:
        return [], [f"Не смог открыть файл: {err}"]

    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], ["Файл пустой."]

    by_title = {_norm(title): key for key, title, _ in COLUMNS}
    header = {}
    for idx, title in enumerate(rows[0]):
        key = by_title.get(_norm(title))
        if key:
            header[key] = idx

    if "fio" not in header:
        return [], ["В файле нет столбца «ФИО» — без него не разобрать. "
                    "Выгрузи актуальный файл кнопкой «Выгрузить» и правь его."]

    parsed, errors = [], []
    seen_ids = set()
    for line, raw in enumerate(rows[1:], start=2):
        def col(key):
            idx = header.get(key)
            return raw[idx] if idx is not None and idx < len(raw) else None

        fio = _to_str(col("fio"))
        if not fio:
            continue  # пустой хвост файла — это норма, молча пропускаем

        item = {"row": line, "fio": fio}
        try:
            raw_id = _to_str(col("id"))
            if raw_id:
                item["id"] = int(float(raw_id))
                if item["id"] in seen_ids:
                    errors.append(f"строка {line}: id {item['id']} встречается второй раз")
                    continue
                seen_ids.add(item["id"])
            else:
                item["id"] = None

            item["birthday"] = _to_birthday(col("birthday"))
            item["is_active"] = _to_bool(col("is_active"))
            studak = _to_str(col("studak"))
            item["studak"] = int(float(studak)) if studak else 0
            for key in TEXT_FIELDS:
                if key == "fio":
                    continue
                item[key] = _to_str(col(key))
            item["tg_username"] = item["tg_username"].lstrip("@")
        except (ValueError, TypeError) as err:
            errors.append(f"строка {line} ({fio}): {err}")
            continue

        parsed.append(item)

    if not parsed and not errors:
        errors.append("В файле не нашлось ни одной строки с ФИО.")
    return parsed, errors
